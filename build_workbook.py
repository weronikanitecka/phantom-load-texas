"""Final data workbook: tab 'Raw data' (only reported/measured values, no calculations) and
tab 'Calculations' (Excel formulas referencing the raw tab). Also exports the raw tab as CSV.
Run: python build_workbook.py ; then recalc ; outputs data/phantom_load_data.xlsx, data/phantom_load_raw_data.csv
"""
from pathlib import Path
import json, sys
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT)); from ercot_hourly import load
HOURLY_URL = 'https://www.ercot.com/gridinfo/load/load_hist/index.html'
ML_URL = 'https://github.com/ml-energy/leaderboard/tree/2bdd0070186956688e1f842dc28077980fe64013/public/data/models'
ZONES = ['COAST','EAST','FWEST','NORTH','NCENT','SOUTH','SCENT','WEST']
BENCH = ['model','task','gpu_model','num_gpus','max_batch_size','energy_per_token_J','avg_output_tokens',
         'avg_power_W','throughput_tokens_per_s','avg_batch_size','median_latency_per_token_ms']
COLS = ['row_id','section','item','date','date_precision','hour_ending','dst_repeat_hour','value','unit'] + ['bench_'+c for c in BENCH] + ['zone_'+z+'_MW' for z in ZONES] + ['source','url']
rows, ref = [], {}
def add(key=None, **kw):
    rows.append(kw)
    if key: ref[key] = len(rows) + 1          # +1 for header row in Excel

# 1-6: reported records (no derived rows)
D = pd.read_csv(ROOT/'data/phantom_load_dataset.csv')
D = D[~D.data_type.str.startswith('derived') & (D.dataset != 'ML.ENERGY Leaderboard v3 benchmark') & (D.dataset != 'Derived values')]
SEC = {'ERCOT large-load queue over time': '1. ERCOT large-load queue (requested load)', 'ERCOT large-load queue by status': '1. ERCOT large-load queue (requested load)',
       'ERCOT energized large load': '2. ERCOT large loads actually running', 'ERCOT Batch Zero screening': '3. ERCOT Batch Zero screening',
       'ERCOT peak demand records': '4. ERCOT peak demand records', 'Policy events': '5. Policy events',
       'AI energy per prompt (company disclosures)': '6. AI energy and usage (company disclosures)', 'AI usage': '6. AI energy and usage (company disclosures)'}
D['sec'] = D.dataset.map(SEC); D = D.sort_values(['sec','dataset','date'])
keymap = {('Queue total','2024-12'):'q1', ('Queue total','2025-11'):'q2', ('Queue total','2025-12-17'):'q3', ('Queue total','2026-03-26'):'q4',
          ('Queue total','2026-05'):'q5', ('Queue total','2026-06-30'):'q6', ('Data-center share of queue','2026-06-30'):'dc',
          ('Observed non-simultaneous monthly peak of approved loads','2025-08'):'r1', ('Observed non-simultaneous monthly peak of approved loads','2025-11'):'r2',
          ('Observed non-simultaneous monthly peak of approved loads','2026-06'):'r3'}
for _, x in D.iterrows():
    k = keymap.get((x['item'], x.date))
    if x['item'].startswith('Batch Zero: base'): k = 'b1'
    if x['item'].startswith('Batch Zero: studied'): k = 'b2'
    if x['item'].startswith('Batch Zero: excluded'): k = 'b3'
    if x['item'].startswith('Gemini median text prompt, full'): k = 'g1'
    if x['item'].startswith('Gemini median text prompt, active'): k = 'g2'
    if x['item'].startswith('Reduction in energy per median Gemini'): k = 'g3'
    if x['item'] == 'ChatGPT average query': k = 'o1'
    if x['item'] == 'ChatGPT prompts per day': k = 'o2'
    val = None if x.unit == 'event' else float(x.value)
    add(k, section=x.sec, date_or_time=x.date, item=x['item'], value=val, unit=None if x.unit == 'event' else x.unit,
        source=f'{x.publisher}: {x.document}', url=x.url)

# 7: all ML.ENERGY benchmark runs
for p in sorted((ROOT/'raw/ml_energy').glob('Qwen__*.json')):
    o = json.loads(p.read_text())
    for c in sorted(o['configurations'], key=lambda c: (c['gpu_model'], c['num_gpus'], c['max_num_seqs'])):
        task = {'gpqa':'GPQA (long reasoning)','lm-arena-chat':'LM Arena (chat)'}[o['task']]
        k = None
        if o['model_id']=='Qwen/Qwen3-32B' and c['gpu_model']=='B200' and c['num_gpus']==1 and c['max_num_seqs']==128:
            k = 'm_reas' if o['task']=='gpqa' else 'm_chat'
        add(k, section='7. ML.ENERGY benchmark runs (measured)', date_or_time='2025-26 (Leaderboard v3)',
            item=f"Energy per response: {o['model_id']}, {task}, {c['gpu_model']} x{c['num_gpus']}, max batch {c['max_num_seqs']}",
            value=c['energy_per_request_joules'], unit='J/response (GPU)', source=f'ML.ENERGY Leaderboard v3: {p.name}', url=ML_URL,
            bench_model=o['model_id'], bench_task=task, bench_gpu_model=c['gpu_model'], bench_num_gpus=c['num_gpus'],
            bench_max_batch_size=c['max_num_seqs'], bench_energy_per_token_J=c['energy_per_token_joules'], bench_avg_output_tokens=c['avg_output_len'],
            bench_avg_power_W=c['avg_power_watts'], bench_throughput_tokens_per_s=c['output_throughput_tokens_per_sec'],
            bench_avg_batch_size=c['avg_batch_size'], bench_median_latency_per_token_ms=c['median_itl_ms'])

# 8: all hourly ERCOT load readings
h = load()
first_hour = len(rows) + 2
for x in h.itertuples(index=False):
    add(section='8. ERCOT hourly load (native load)', date_or_time=x[0], item='Hourly native load, ERCOT total',
        value=float(x.ERCOT), unit='MW', source=f'ERCOT Hourly Load Data Archive: {x.source_file}', url=HOURLY_URL,
        **{'zone_'+z+'_MW': float(getattr(x, z)) for z in ZONES})
last_hour = len(rows) + 1

raw = pd.DataFrame(rows)
def parse(t):
    t = str(t)
    if '/' in t:                       # ERCOT hourly: 'MM/DD/YYYY HH:MM' (+ ' DST' for the repeated hour)
        return pd.Timestamp(pd.to_datetime(t[:10], format='%m/%d/%Y')), 'hour', int(t[11:13]), t.endswith('DST')
    if len(t) == 10: return pd.Timestamp(t), 'day', None, None
    if len(t) == 7 and t[4] == '-' and t[5] != 'Q': return pd.Timestamp(t + '-01'), 'month', None, None
    if '-Q' in t: return pd.Timestamp(f"{t[:4]}-{3*int(t[-1])-2:02d}-01"), 'quarter', None, None
    return pd.NaT, 'not dated (benchmark release 2025-26)', None, None
P = raw.date_or_time.map(parse)
raw['date'] = [x[0] for x in P]; raw['date_precision'] = [x[1] for x in P]
raw['hour_ending'] = pd.array([x[2] for x in P], dtype='Int64'); raw['dst_repeat_hour'] = [None if x[3] is None else ('yes' if x[3] else 'no') for x in P]
raw = raw.reindex(columns=COLS); raw['row_id'] = range(1, len(raw)+1)
csv = raw.copy(); csv['date'] = csv.date.dt.strftime('%Y-%m-%d')
csv.to_csv(ROOT/'data/phantom_load_raw_data.csv', index=False)

# ---- workbook ---------------------------------------------------------------------------
wb = Workbook(); ws = wb.active; ws.title = 'Raw data'
ws.append(COLS)
for r in raw.itertuples(index=False):
    ws.append([None if (v is None or v is pd.NA or (isinstance(v, float) and pd.isna(v)) or v is pd.NaT) else (v.to_pydatetime() if isinstance(v, pd.Timestamp) else v) for v in r])
for c in ws[1]:
    c.font = Font(name='Arial', bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1A1A1A')
for i, w in enumerate([7, 36, 60, 11, 14, 10, 10, 14, 16] + [14]*11 + [11]*8 + [50, 40]):
    ws.column_dimensions[ws.cell(1, i+1).column_letter].width = w
for r in range(2, len(raw)+2): ws.cell(r, 4).number_format = 'yyyy-mm-dd'
ws.freeze_panes = 'A2'

cs = wb.create_sheet('Calculations'); R = "'Raw data'"
F = lambda key, col='H': f"{R}!${col}${ref[key]}"
cs.append(['Calculations. Every value is an Excel formula that points to cells in the "Raw data" tab.'])
cs.append([])
cs.append(['A. Key numbers', None, None, None, None])
cs.append(['Item', 'Value', 'Unit', 'Used in', 'How it is calculated'])
calc = [
 ('Requested load, Dec 2024', f"={F('q1')}/1000", 'GW', 'Chart (orange line)', 'MW / 1000'),
 ('Requested load, Nov 2025', f"={F('q2')}/1000", 'GW', 'Chart (orange line)', 'MW / 1000'),
 ('Requested load, 17 Dec 2025', f"={F('q3')}/1000", 'GW', 'Chart (orange line)', 'MW / 1000'),
 ('Requested load, 26 Mar 2026', f"={F('q4')}/1000", 'GW', 'Chart (orange line)', 'MW / 1000'),
 ('Requested load, May 2026', f"={F('q5')}/1000", 'GW', 'Chart (orange line)', 'MW / 1000'),
 ('Requested load, Jun 2026', f"={F('q6')}/1000", 'GW', 'Chart (orange line), takeaway', 'MW / 1000'),
 ('Large loads actually running, Aug 2025', f"={F('r1')}/1000", 'GW', 'Chart (blue line)', 'MW / 1000'),
 ('Large loads actually running, Nov 2025', f"={F('r2')}/1000", 'GW', 'Chart (blue line)', 'MW / 1000'),
 ('Large loads actually running, Jun 2026', f"={F('r3')}/1000", 'GW', 'Chart (blue line), takeaway', 'MW / 1000'),
 ('Batch Zero: load that passed first screening', f"=({F('b1')}+{F('b2')})/1000", 'GW', 'Chart, voiceover', 'Base load + studied load'),
 ('Batch Zero: load excluded', f"={F('b3')}/1000", 'GW', 'Chart', 'MW / 1000'),
 ('Growth in requested load, Dec 2024 to Jun 2026', f"={F('q6')}/{F('q1')}", 'x', 'Context', 'Jun 2026 / Dec 2024'),
 ('Increase in requested load, Dec 2024 to Jun 2026', f"=({F('q6')}-{F('q1')})/1000", 'GW', 'Voiceover', 'Jun 2026 minus Dec 2024'),
]
start = cs.max_row + 1
for c in calc: cs.append(list(c))
cref = {name: f'B{start+i}' for i, (name, *_ ) in enumerate(calc)}
# monthly table first needed for peaks: place it below key numbers; reference by row
months = sorted(h.month.unique())
key2 = [
 ('Peak hourly demand, 2024', 'PEAK2024', 'GW', 'Takeaway, voiceover', 'Highest monthly peak in 2024 (section B)'),
 ('Peak hourly demand, 2026 (Jan to Sep)', 'PEAK2026', 'GW', 'Chart label, takeaway, voiceover', 'Highest monthly peak in 2026 (section B)'),
 ('Increase in actual peak demand, 2024 to 2026', 'INC', 'GW', 'Takeaway, voiceover', '2026 peak minus 2024 peak'),
 ('Latest requested load vs record peak demand', 'MULT', 'x', 'Takeaway', 'Jun 2026 requested load / 2026 peak'),
 ('Reasoning vs chat energy per response (Qwen3-32B, B200)', f"={F('m_reas')}/{F('m_chat')}", 'x', 'Voiceover (~17x)', 'GPQA run energy / chat run energy'),
 ('Overhead factor: full data center vs chips only', f"={F('g1')}/{F('g2')}", 'x', 'Input', 'Google 0.24 Wh / 0.10 Wh'),
 ('Long reasoning answer, full data center', f"={F('m_reas')}/3600*{F('g1')}/{F('g2')}", 'Wh', 'Input', 'J / 3600 x overhead factor'),
 ('Power for all ChatGPT prompts at 0.34 Wh', f"={F('o2')}*{F('o1')}/24/1E9", 'GW', 'Voiceover', 'prompts/day x Wh / 24 h / 1e9'),
 ('Power for all ChatGPT prompts if all were long reasoning', 'CGHI', 'GW', 'Voiceover', 'prompts/day x reasoning Wh / 24 h / 1e9'),
 ('Long reasoning answers per day that 1 GW can run', 'PERGW', 'answers/day', 'Voiceover (~16 billion)', '1e9 W x 24 h / reasoning Wh'),
 ('1 GW as multiples of ChatGPT daily traffic', 'XCG', 'x', 'Voiceover (>6x)', 'answers per GW / ChatGPT prompts per day'),
 ('Requested load (Jun 2026) as multiples of ChatGPT', 'REQCG', 'x', 'Voiceover (~3,000)', 'requested GW x answers per GW / ChatGPT prompts per day'),
]
s2 = cs.max_row + 1
for c in key2: cs.append(list(c))
k2 = {name: f'B{s2+i}' for i, (name, *_ ) in enumerate(key2)}
cs.append([]); cs.append(['B. Monthly Texas demand from hourly data (grey line = monthly peak)'])
cs.append(['Month', 'Month start', 'Hours', 'Peak hourly demand (MW)', 'Average demand (MW)', 'Peak (GW)'])
m0 = cs.max_row + 1
HR = lambda col: f"{R}!${col}${first_hour}:${col}${last_hour}"
for i, m in enumerate(months):
    r = m0 + i
    cs.append([m, pd.Timestamp(m + '-01').to_pydatetime(),
               f"=COUNTIFS({HR('D')},\">=\"&B{r},{HR('D')},\"<\"&EDATE(B{r},1))",
               f"=_xlfn.MAXIFS({HR('H')},{HR('D')},\">=\"&B{r},{HR('D')},\"<\"&EDATE(B{r},1))",
               f"=AVERAGEIFS({HR('H')},{HR('D')},\">=\"&B{r},{HR('D')},\"<\"&EDATE(B{r},1))",
               f"=D{r}/1000"])
    cs[f'B{r}'].number_format = 'yyyy-mm-dd'
    for col, fmt in zip('CDEF', ['#,##0', '#,##0', '#,##0', '0.0']): cs[f'{col}{r}'].number_format = fmt
m1 = cs.max_row
mrows = {m: m0+i for i, m in enumerate(months)}
y24 = [mrows[m] for m in months if m.startswith('2024')]; y26 = [mrows[m] for m in months if m.startswith('2026')]
cs[k2['Peak hourly demand, 2024']] = f"=MAX(F{y24[0]}:F{y24[-1]})"
cs[k2['Peak hourly demand, 2026 (Jan to Sep)']] = f"=MAX(F{y26[0]}:F{y26[-1]})"
cs[k2['Increase in actual peak demand, 2024 to 2026']] = f"={k2['Peak hourly demand, 2026 (Jan to Sep)']}-{k2['Peak hourly demand, 2024']}"
cs[k2['Latest requested load vs record peak demand']] = f"={cref['Requested load, Jun 2026']}/{k2['Peak hourly demand, 2026 (Jan to Sep)']}"
rw = k2['Long reasoning answer, full data center']
cs[k2['Power for all ChatGPT prompts if all were long reasoning']] = f"={F('o2')}*{rw}/24/1E9"
cs[k2['Long reasoning answers per day that 1 GW can run']] = f"=1E9*24/{rw}"
cs[k2['1 GW as multiples of ChatGPT daily traffic']] = f"={k2['Long reasoning answers per day that 1 GW can run']}/{F('o2')}"
cs[k2['Requested load (Jun 2026) as multiples of ChatGPT']] = f"={cref['Requested load, Jun 2026']}*{k2['1 GW as multiples of ChatGPT daily traffic']}"
for r in range(start, s2 + len(key2)):
    unit = cs[f'C{r}'].value
    cs[f'B{r}'].number_format = {'GW': '0.000', 'x': '0.0', 'Wh': '0.000', 'answers/day': '#,##0'}.get(unit, 'General')
# style
for r in (1,): cs[f'A{r}'].font = Font(name='Arial', italic=True, color='555555')
for r in (3, m0-2): cs[f'A{r}'].font = Font(name='Arial', bold=True, size=12)
for r in (4, m0-1):
    for c in cs[r]: c.font = Font(name='Arial', bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1A1A1A')
for col, w in zip('ABCDEF', [52, 18, 12, 30, 52, 10]): cs.column_dimensions[col].width = w
wb.save(ROOT/'data/phantom_load_data.xlsx')
print('raw rows', len(raw), '| hourly rows', first_hour, '-', last_hour, '| refs', ref)
